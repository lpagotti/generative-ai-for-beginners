"""Exporta o cadastro de clientes do MOGO Gourmet (Cadastros > Pessoas > Clientes) para Excel.

Login com as variaveis de ambiente MOGO_USER e MOGO_PASS. A listagem usa a mesma
chamada interna da tela (POST /Cadastros/GenerioDadosJqGrid), paginada.

Uso:
    python exportar_clientes.py [--empresa cidinhabolos] [--saida clientes.xlsx] [--por-pagina 500]
"""

import argparse
import os
import sys
import time
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright

BASE_URL = "https://app3.mogogourmet.com.br"
TIPO_CLIENTE = "kardapio.Suprimentos.Cliente"

# (cabecalho no Excel, campo na API)
COLUNAS = [
    ("Id", "Id"),
    ("Nome", "Nome"),
    ("Tipo pessoa", "TipoPessoa"),
    ("CPF/CNPJ", "CNPJ_CPF"),
    ("Razão social", "RazaoSocial"),
    ("Fantasia", "Fantasia"),
    ("Inscrição estadual", "InscricaoEstadual"),
    ("Celular", "TelefoneCelular"),
    ("Telefone residencial", "TelefoneResidencial"),
    ("Telefone comercial", "TelefoneComercial"),
    ("E-mail", "Email"),
    ("CEP", "CEP"),
    ("Logradouro", "Logradouro"),
    ("Número", "Numero"),
    ("Complemento", "Complemento"),
    ("Bairro", "Bairro"),
    ("Cidade", "Cidade_Nome"),
    ("UF", "UFCidade_Descricao"),
    ("Sexo", "Sexo"),
    ("Nascimento", "Nascimento"),
    ("Data cadastro", "DataCadastro"),
    ("Origem cadastro", "OrigemCadastro"),
    ("Qtd visitas", "QuantidadeVisitas"),
    ("Qtd delivery", "QuantidadeDelivery"),
    ("Valor consumido", "ValorConsumido"),
    ("Último pedido", "DataUltimoPedido"),
    ("Observações", "ObservacoesCliente"),
    ("Ativo", "Ativo"),
]

FETCH_JS = """async form => {
    const r = await fetch('/Cadastros/GenerioDadosJqGrid', {
        method: 'POST',
        body: new URLSearchParams(form),
        headers: {'X-Requested-With': 'XMLHttpRequest'},
    });
    if (!r.ok) throw new Error('HTTP ' + r.status);
    return await r.json();
}"""


def login(page, empresa, usuario, senha):
    page.goto(f"{BASE_URL}/{empresa}", wait_until="networkidle", timeout=60000)
    form = page.locator("form").filter(has=page.locator("#button-login"))
    form.locator("input[name=UserName]").fill(usuario)
    form.locator("input[name=Password]").fill(senha)
    page.click("#button-login")
    page.wait_for_load_state("networkidle", timeout=60000)

    # Tela "Escolha o modo de acesso"
    gerencial = page.get_by_text("Gerencial", exact=True)
    if gerencial.count() == 0:
        raise RuntimeError("Login falhou: confira MOGO_USER / MOGO_PASS.")
    gerencial.first.click()
    # O MOGO permite uma sessao por usuario; se houver outra aberta, encerra
    finalizar = page.get_by_text("Finalizar sessão ativa").locator("visible=true")
    if finalizar.count():
        print("Usuario ja conectado em outro dispositivo - finalizando a sessao anterior.")
        finalizar.click()
    page.locator("button:visible", has_text="Entrar").last.click()
    try:
        page.wait_for_url(lambda url: "/Account/" not in url, timeout=60000)
    except Exception:
        raise RuntimeError(f"Login nao concluido (parou em {page.url}).") from None
    page.wait_for_load_state("networkidle", timeout=60000)


def buscar_pagina(page, numero, por_pagina):
    form = {
        "filtro": "", "cTipo": TIPO_CLIENTE, "idEstab": "0",
        "todosEstabelecimentos": "false", "filtrosEnv": "", "Searching": "true",
        "rows": str(por_pagina), "PageIndex": str(numero), "page": str(numero),
        "SortingName": "", "sord": "asc", "sidx": "Id", "_search": "true",
    }
    for tentativa in range(3):
        try:
            return page.evaluate(FETCH_JS, form)
        except Exception as erro:
            if tentativa == 2:
                raise
            print(f"  pagina {numero}: {erro} - tentando de novo", file=sys.stderr)
            time.sleep(2 * (tentativa + 1))


def valor(cliente, campo):
    v = cliente.get(campo)
    if isinstance(v, bool):
        return "Sim" if v else "Não"
    if isinstance(v, (dict, list)):
        return None
    return v


def salvar_excel(clientes, caminho):
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    ws.append([cab for cab, _ in COLUNAS])
    for c in clientes:
        ws.append([valor(c, campo) for _, campo in COLUNAS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, (cab, campo) in enumerate(COLUNAS, start=1):
        maior = max([len(cab)] + [len(str(valor(c, campo) or "")) for c in clientes[:2000]])
        ws.column_dimensions[get_column_letter(i)].width = min(maior + 2, 50)
    wb.save(caminho)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--empresa", default="cidinhabolos", help="dominio da empresa no MOGO")
    parser.add_argument("--saida", default=None, help="arquivo .xlsx de saida")
    parser.add_argument("--por-pagina", type=int, default=500, help="registros por chamada")
    args = parser.parse_args()

    usuario, senha = os.environ.get("MOGO_USER"), os.environ.get("MOGO_PASS")
    if not usuario or not senha:
        sys.exit("Defina as variaveis de ambiente MOGO_USER e MOGO_PASS.")

    saida = Path(args.saida or f"saida/clientes_{args.empresa}_{datetime.now():%Y%m%d_%H%M}.xlsx")
    saida.parent.mkdir(parents=True, exist_ok=True)

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        print(f"Entrando no MOGO ({args.empresa})...")
        login(page, args.empresa, usuario, senha)

        primeira = buscar_pagina(page, 1, args.por_pagina)
        total, paginas = primeira["records"], primeira["total"]
        print(f"{total} clientes em {paginas} paginas de {args.por_pagina}.")
        clientes = list(primeira["rows"])
        try:
            for n in range(2, paginas + 1):
                clientes.extend(buscar_pagina(page, n, args.por_pagina)["rows"])
                print(f"  pagina {n}/{paginas} - {len(clientes)} registros")
        finally:
            # Libera a sessao para nao bloquear o proximo login
            page.goto(f"{BASE_URL}/Account/LogOff")
            browser.close()

    unicos = {c["Id"]: c for c in clientes}
    clientes = sorted(unicos.values(), key=lambda c: int(c["Id"]))
    if len(clientes) != total:
        print(f"AVISO: esperados {total}, recebidos {len(clientes)} unicos.", file=sys.stderr)

    salvar_excel(clientes, saida)
    print(f"OK: {len(clientes)} clientes salvos em {saida.resolve()}")


if __name__ == "__main__":
    main()
